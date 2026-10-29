"""Local bounded macro-free OOXML editing; never launches Office or external links."""
from __future__ import annotations

import io
import math
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree


def validate(path, data, *, workspace=None):
    from .github_safety import scan_text
    if Path(path).suffix.lower() not in {'.docx', '.xlsx', '.pptx'} or len(data) > 500_000:
        raise ValueError('Office supports bounded .docx/.xlsx/.pptx only')
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 200 or sum(e.file_size for e in entries) > 5_000_000:
            raise ValueError('Office archive expansion budget exceeded')
        names = [e.filename for e in entries]
        if len(names) != len(set(names)) or '[Content_Types].xml' not in names:
            raise ValueError('Invalid Office archive')
        required = {'.docx': 'word/document.xml', '.xlsx': 'xl/workbook.xml', '.pptx': 'ppt/presentation.xml'}[Path(path).suffix.lower()]
        if required not in names:
            raise ValueError('Office extension/content mismatch')
        for entry in entries:
            name = entry.filename
            if (name.startswith('/') or '..' in Path(name).parts or '\\' in name or entry.flag_bits & 1
                    or any(word in name.lower() for word in ('vbaproject', 'embeddings/', 'externallinks/', 'activex/', '.bin'))):
                raise ValueError('Unsafe Office member/macro/embedded content')
            content = archive.read(entry)
            if name.lower().endswith(('.xml', '.rels')):
                text = content.decode('utf-8')
                if '<!DOCTYPE' in text.upper() or '<!ENTITY' in text.upper():
                    raise ValueError('Office XML entities are refused')
                root = ElementTree.fromstring(text)
                if any(node.attrib.get('TargetMode', '').lower() == 'external' or node.tag.rsplit('}', 1)[-1] == 'f'
                       for node in root.iter()):
                    raise ValueError('Office formulas/external relationships are refused')
                scan_text(text, workspace=workspace)


def operate(path, args):
    action = args['action']
    output = io.BytesIO()
    suffix = path.suffix.lower()
    data = path.read_bytes() if path.exists() else None
    if data is not None:
        validate(path, data)
    elif action == 'read':
        raise ValueError('Office document does not exist')
    if suffix == '.docx':
        from docx import Document
        document = Document(io.BytesIO(data)) if data else Document()
        if action == 'read':
            return '\n'.join(p.text for p in document.paragraphs)[:32_000], None
        if action not in {'write', 'append'}:
            raise ValueError('DOCX actions: read/write/append')
        text = args.get('text')
        if not isinstance(text, str) or len(text) > 32_000:
            raise ValueError('Bounded paragraph text required')
        if action == 'write':
            document = Document()
        for paragraph in text.splitlines():
            document.add_paragraph(paragraph)
        document.save(output)
    elif suffix == '.xlsx':
        import openpyxl
        book = openpyxl.load_workbook(io.BytesIO(data), keep_links=False) if data else openpyxl.Workbook()
        sheet = args.get('sheet') or book.sheetnames[0]
        if not isinstance(sheet, str) or sheet not in book.sheetnames:
            raise ValueError('Select an existing sheet')
        cell = args.get('cell', 'A1')
        if not isinstance(cell, str) or not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]{0,6}', cell):
            raise ValueError('A bounded cell coordinate is required')
        from openpyxl.utils.cell import column_index_from_string, coordinate_from_string
        column, row = coordinate_from_string(cell)
        if row > 1048576 or column_index_from_string(column) > 16384:
            raise ValueError('Cell outside Excel limits')
        if action == 'read':
            return str(book[sheet][cell].value)[:32_000], None
        if action != 'set':
            raise ValueError('XLSX actions: read/set')
        value = args.get('value')
        if (type(value) not in {str, int, float, bool, type(None)} or (type(value) is float and not math.isfinite(value))
                or isinstance(value, str) and (len(value) > 10_000 or value.lstrip().startswith(('=', '+', '@')))):
            raise ValueError('Literal bounded cell value required; formulas are refused')
        book[sheet][cell] = value
        book.save(output)
    elif suffix == '.pptx':
        from pptx import Presentation
        from pptx.util import Inches
        deck = Presentation(io.BytesIO(data)) if data else Presentation()
        if action == 'read':
            return '\n'.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)[:32_000], None
        if action not in {'write', 'append'} or not isinstance(args.get('text'), str) or len(args['text']) > 32_000:
            raise ValueError('PPTX supports bounded write/append slide text')
        if action == 'write':
            deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(5)).text = args['text']
        for relation in list(deck.part.rels.values()):
            if relation.reltype.endswith('/printerSettings'):
                deck.part.drop_rel(relation.rId)
        deck.save(output)
    else:
        raise ValueError('Only .docx/.xlsx/.pptx supported')
    planned = output.getvalue()
    validate(path, planned)
    return 'Local Office document updated; no external app or link executed', planned
